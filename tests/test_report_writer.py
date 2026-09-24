import openpyxl
import pandas as pd
from src.report_writer import export_report


def test_export_report_sheets(tmp_path):
    df = pd.DataFrame({
        "status": ["inativo", "ativo"],
        "economia_anual_inativo": [420.0, 0.0]
    })
    out_file = tmp_path / "relatorio.xlsx"
    export_report(df, str(out_file))

    assert out_file.exists()
    wb = openpyxl.load_workbook(out_file)
    assert "Resumo" in wb.sheetnames
    assert "Detalhes" in wb.sheetnames
